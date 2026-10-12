"""
Excel con formulas REALES de Excel para la red TS-LIF.
Cambia un parametro en la hoja Parametros y toda la simulacion recalcula sola.
"""
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference
import os

BASE = os.path.dirname(os.path.abspath(__file__))
CSV  = os.path.join(BASE, 'subirTarea', 'electricity_sample.csv')
OUT  = os.path.join(BASE, 'TS_LIF_calculo_paso_a_paso.xlsx')

data   = np.loadtxt(CSV, delimiter=',')
N_ROWS = len(data)   # 168 pasos

wb = openpyxl.Workbook()

# ─── helpers de estilo ───────────────────────────────────────────────────────
def fill(h):
    return PatternFill('solid', fgColor=h if len(h)==8 else 'FF'+h)

def brd(thick_right=False):
    t = Side(style='thin',   color='CCCCCC')
    m = Side(style='medium', color='555555')
    r = m if thick_right else t
    return Border(left=t, right=r, top=t, bottom=t)

def put(ws, row, col, value=None, formula=None, bg=None, fg='000000',
        bold=False, size=10, align='center', wrap=False, fmt=None,
        italic=False, thick_right=False):
    c = ws.cell(row=row, column=col)
    c.value = formula if formula else value
    c.font = Font(bold=bold, color='FF'+fg, size=size, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
    if bg:
        c.fill = fill(bg)
    c.border = brd(thick_right)
    if fmt:
        c.number_format = fmt
    return c

# Colores
C = {
    'az_osc': '1F3864', 'az_clr': 'D6E4F0',
    'vd_osc': '1E5631', 'vd_clr': 'D0E8D0', 'vd_brt': '90EE90',
    'rj_osc': '7B0000', 'rj_clr': 'FCE4D6', 'rj_brt': 'FFB3B3',
    'mr_osc': '4A235A', 'mr_clr': 'E8DAEF',
    'gr_osc': '404040', 'gr_clr': 'F0F0F0',
    'am':     'FFFF99', 'blanco': 'FFFFFF', 'naranja': 'FFE0B2',
}

# =============================================================================
# HOJA 1 — PARAMETROS
# Las celdas B4:B11 son las que usan todas las formulas de Simulacion.
# NUNCA mover esas celdas.
# =============================================================================
ws_p = wb.active
ws_p.title = 'Parametros'
ws_p.column_dimensions['A'].width = 22
ws_p.column_dimensions['B'].width = 12
ws_p.column_dimensions['C'].width = 16
ws_p.column_dimensions['D'].width = 55

# Fila 1: titulo
ws_p.merge_cells('A1:D1')
c = ws_p.cell(row=1, column=1,
    value='PARAMETROS — Edita la columna B y toda la simulacion se actualiza')
c.font = Font(bold=True, size=13, color='FFFFFFFF')
c.fill = fill(C['az_osc'])
c.alignment = Alignment(horizontal='center', vertical='center')
ws_p.row_dimensions[1].height = 28

# Fila 2: aviso
ws_p.merge_cells('A2:D2')
c = ws_p.cell(row=2, column=1,
    value='Las celdas B4:B11 son referenciadas por formulas en la hoja Simulacion. '
          'No mover. Cambiar el valor y Excel recalcula todo automaticamente.')
c.font = Font(italic=True, size=10)
c.fill = fill('FFFFF3CD')
c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws_p.row_dimensions[2].height = 24

# Fila 3: cabeceras
for col, txt in [(1,'Parametro'),(2,'VALOR'),(3,'Rango valido'),(4,'Que controla')]:
    put(ws_p, 3, col, txt, bg=C['gr_osc'], fg='FFFFFF', bold=True)

# Filas 4-11: los 8 parametros
# IMPORTANTE: filas 4-11 = B4:B11 referenciadas en Simulacion
PARAMS = [                                    # fila en Parametros
    (4,  'a1  (alpha1)',  0.95, '0.01 – 0.99',
     'Decaimiento de la DENDRITA. Cerca de 1 = olvida despacio = memoria larga → paso-bajo'),
    (5,  'a2  (alpha2)',  0.05, '0.01 – 0.99',
     'Decaimiento del SOMA. Cerca de 0 = olvida rapido = memoria corta → paso-alto'),
    (6,  'b1  (beta1)',   0.00, '-1.0 – 1.0',
     'Acoplamiento soma→dendrita. VALE 0 POR DISENO DEL PAPER: '
     'el soma no retroalimenta a la dendrita. La columna "b1·vs_prev" en Simulacion '
     'siempre sera 0 mientras b1=0. Prueba cambiarlo a 0.1 para ver el efecto.'),
    (7,  'b2  (beta2)',  -0.90, '-1.0 – 0.0',
     'Acoplamiento dendrita→soma. Negativo = la dendrita frena al soma (evita saturacion)'),
    (8,  'g1  (gamma1)', 0.50, '0.0 – 2.0',
     'Reset dendrita: vd baja en g1 cada vez que sd_prev=1 (la neurona disparo antes)'),
    (9,  'g2  (gamma2)', 1.00, '0.0 – 2.0',
     'Reset soma: vs baja en g2 cada vez que ss_prev=1'),
    (10, 'v_th',          0.50, '0.1 – 1.0',
     'Umbral de disparo: sd=1 si vd >= v_th  /  ss=1 si vs >= v_th  '
     '(0.5 para datos normalizados [0,1] — con 1.0 la neurona nunca dispara)'),
    (11, 'kappa',         0.50, '0.0 – 1.0',
     'Peso de mezcla salida: s_mix = kappa*sd + (1-kappa)*ss'),
]

for p_row, nombre, valor, rango, desc in PARAMS:
    bg = C['az_clr'] if p_row % 2 == 0 else C['blanco']
    put(ws_p, p_row, 1, nombre, bg=bg, align='left', bold=True)
    c = ws_p.cell(row=p_row, column=2, value=valor)
    c.font = Font(bold=True, size=13, color='FF1F3864')
    c.fill = fill(C['naranja'])
    c.border = Border(left=Side(style='medium', color='555555'),
                      right=Side(style='medium', color='555555'),
                      top=Side(style='medium', color='555555'),
                      bottom=Side(style='medium', color='555555'))
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.number_format = '0.00'
    put(ws_p, p_row, 3, rango, bg=bg, italic=True, fg='555555')
    put(ws_p, p_row, 4, desc,  bg=bg, align='left', wrap=True)
    ws_p.row_dimensions[p_row].height = 22

# Fila 12: espacio
ws_p.row_dimensions[12].height = 10

# Fila 13: titulo estabilidad
ws_p.merge_cells('A13:D13')
c = ws_p.cell(row=13, column=1,
    value='VERIFICACION DE ESTABILIDAD (Teorema 1) — formulas Excel que recalculan con los parametros')
c.font = Font(bold=True, size=11, color='FFFFFFFF')
c.fill = fill(C['gr_osc'])
c.alignment = Alignment(horizontal='center', vertical='center')
ws_p.row_dimensions[13].height = 22

# Fila 14: cabecera tabla estabilidad
for col, txt in [(1,'Magnitud'),(2,'Resultado'),(3,'Formula'),(4,'Que significa')]:
    put(ws_p, 14, col, txt, bg='D9D9D9', bold=True)

# Filas 15-20: calculos de estabilidad con formulas Excel (en ingles para openpyxl)
# B4=a1, B5=a2, B6=b1, B7=b2
ESTAB = [
    # (fila, etiqueta, formula_excel, descripcion)
    (15, 'Traza',
     '=B4+B5+B6*B7',
     'a1 + a2 + b1*b2 — suma de los terminos de decaimiento mas el producto de acoplamientos'),
    (16, 'Discriminante',
     '=SQRT(B15^2-4*B4*B5)',
     'sqrt(traza^2 - 4*a1*a2) — parte variable bajo la raiz del polinomio caracteristico'),
    (17, 'Lambda 1  (dendrita)',
     '=(B15+B16)/2',
     'Autovalor 1 = (traza + disc)/2 — constante de tiempo de la DENDRITA (debe ser |λ|<1)'),
    (18, 'Lambda 2  (soma)',
     '=(B15-B16)/2',
     'Autovalor 2 = (traza - disc)/2 — constante de tiempo del SOMA (debe ser |λ|<1)'),
    (19, '|Lambda 1| < 1 ?',
     '=IF(ABS(B17)<1,"ESTABLE","INESTABLE — ajusta los parametros")',
     'Si es INESTABLE, el potencial crecera sin limite y la red divergira'),
    (20, '|Lambda 2| < 1 ?',
     '=IF(ABS(B18)<1,"ESTABLE","INESTABLE — ajusta los parametros")',
     'Ambos autovalores deben cumplir esta condicion para garantizar convergencia'),
]

for e_row, etiq, formula, desc in ESTAB:
    bg = C['gr_clr'] if e_row % 2 == 0 else C['blanco']
    put(ws_p, e_row, 1, etiq,    bg=bg, align='left', bold=True)
    c = ws_p.cell(row=e_row, column=2)
    c.value = formula
    c.font = Font(bold=True, size=12, color='FF1F3864')
    c.fill = fill(C['naranja'])
    c.border = Border(left=Side(style='medium', color='555555'),
                      right=Side(style='medium', color='555555'),
                      top=Side(style='medium', color='555555'),
                      bottom=Side(style='medium', color='555555'))
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.number_format = '0.0000'
    put(ws_p, e_row, 3, formula, bg=bg, align='left', size=9, italic=True, fg='555555')
    put(ws_p, e_row, 4, desc,    bg=bg, align='left', wrap=True)
    ws_p.row_dimensions[e_row].height = 22

ws_p.freeze_panes = 'A4'

# =============================================================================
# HOJA 2 — DATOS (electricity_sample.csv)
# Los valores de c[t] vienen de aqui. Simulacion los referencia.
# c[t=0] esta en Datos!B3, c[t=1] en Datos!B4, etc.
# =============================================================================
ws_d = wb.create_sheet('Datos')
ws_d.column_dimensions['A'].width = 8
for col in ['B','C','D']:
    ws_d.column_dimensions[col].width = 16

ws_d.merge_cells('A1:D1')
c = ws_d.cell(row=1, column=1,
    value='electricity_sample.csv — 168 horas x 3 nodos reales UCI Electricity dataset')
c.font = Font(bold=True, size=12, color='FFFFFFFF')
c.fill = fill(C['az_osc'])
c.alignment = Alignment(horizontal='center', vertical='center')
ws_d.row_dimensions[1].height = 24

for col, txt in [(1,'Hora t'),(2,'Nodo 0'),(3,'Nodo 1'),(4,'Nodo 2')]:
    put(ws_d, 2, col, txt, bg=C['gr_osc'], fg='FFFFFF', bold=True)
ws_d.row_dimensions[2].height = 20

# DATA_ROW_START = 3  (c[t=0] en fila 3, c[t=k] en fila 3+k)
DATA_ROW_START = 3
for t in range(N_ROWS):
    row = DATA_ROW_START + t
    bg = C['gr_clr'] if t % 2 == 0 else C['blanco']
    put(ws_d, row, 1, t,         bg=C['am'],  bold=True, fmt='0')
    put(ws_d, row, 2, data[t,0], bg=bg, fmt='0.000000')
    put(ws_d, row, 3, data[t,1], bg=bg, fmt='0.000000')
    put(ws_d, row, 4, data[t,2], bg=bg, fmt='0.000000')

ws_d.freeze_panes = 'A3'

# =============================================================================
# HOJA 3 — SIMULACION CON FORMULAS EXCEL REALES
#
# Referencias a Parametros (NUNCA cambiar estas filas):
#   a1  = Parametros!$B$4
#   a2  = Parametros!$B$5
#   b1  = Parametros!$B$6
#   b2  = Parametros!$B$7
#   g1  = Parametros!$B$8
#   g2  = Parametros!$B$9
#   v_th= Parametros!$B$10
#   kap = Parametros!$B$11
#
# Layout de columnas:
#   A=t  B=c[t]  C=a1*vd_prev  D=b1*vs_prev  E=(1-a1)*c  F=-g1*sd_prev
#   G=vd[t]  H=sd[t]  | sep |  J=a2*vs_prev  K=b2*vd[t]
#   L=(1-a2)*c  M=-g2*ss_prev  N=vs[t]  O=ss[t]  | sep |  Q=s_mix[t]
#
# Fila INIT_ROW (estado inicial): vd=0, vs=0, sd=0, ss=0
# Filas INIT_ROW+1 ... INIT_ROW+N_ROWS: formulas reales
# =============================================================================
ws = wb.create_sheet('Simulacion')

# Referencias a Parametros (strings para usar en formulas)
PA1  = 'Parametros!$B$4'
PA2  = 'Parametros!$B$5'
PB1  = 'Parametros!$B$6'
PB2  = 'Parametros!$B$7'
PG1  = 'Parametros!$B$8'
PG2  = 'Parametros!$B$9'
PVTH = 'Parametros!$B$10'
PKAP = 'Parametros!$B$11'

# Columnas fijas
COL = {
    't':    1,   # A
    'c':    2,   # B
    't1d':  3,   # C  a1*vd_prev
    't2d':  4,   # D  b1*vs_prev
    't3d':  5,   # E  (1-a1)*c
    't4d':  6,   # F  -g1*sd_prev
    'vd':   7,   # G
    'sd':   8,   # H
    'sep1': 9,   # I  separador
    't1s':  10,  # J  a2*vs_prev
    't2s':  11,  # K  b2*vd[t]
    't3s':  12,  # L  (1-a2)*c
    't4s':  13,  # M  -g2*ss_prev
    'vs':   14,  # N
    'ss':   15,  # O
    'sep2': 16,  # P  separador
    'mix':  17,  # Q
}

# Letras de columna (para formulas)
L = {k: get_column_letter(v) for k, v in COL.items()}

# Anchos
for col, w in [(1,5),(2,10),(3,14),(4,14),(5,13),(6,14),(7,11),(8,8),
               (9,2),(10,14),(11,14),(12,13),(13,14),(14,11),(15,8),
               (16,2),(17,11)]:
    ws.column_dimensions[get_column_letter(col)].width = w

# Fila 1: titulo
ws.merge_cells('A1:Q1')
c = ws.cell(row=1, column=1,
    value='SIMULACION TS-LIF — Formulas Excel reales. '
          'Cambia Parametros!B4:B11 y todo recalcula automaticamente.')
c.font = Font(bold=True, size=13, color='FFFFFFFF')
c.fill = fill(C['az_osc'])
c.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 28

# Fila 2: grupos
ws.merge_cells('A2:B2')
put(ws, 2, 1, 'ENTRADA', bg=C['gr_osc'], fg='FFFFFF', bold=True)
ws.merge_cells('C2:H2')
put(ws, 2, 3,
    'DENDRITA  →  vd[t] = a1·vd_prev  +  b1·vs_prev  +  (1−a1)·c[t]  −  g1·sd_prev',
    bg=C['vd_osc'], fg='FFFFFF', bold=True)
ws.merge_cells('J2:P2')
put(ws, 2, 10,
    'SOMA  →  vs[t] = a2·vs_prev  +  b2·vd[t]  +  (1−a2)·c[t]  −  g2·ss_prev',
    bg=C['rj_osc'], fg='FFFFFF', bold=True)
put(ws, 2, 17, 'SALIDA', bg=C['mr_osc'], fg='FFFFFF', bold=True)
ws.row_dimensions[2].height = 22

# Fila 3: cabeceras con nombre de variable y formula en texto
HEADS = [
    (COL['t'],   't\nhora',                    C['gr_osc'], 'FFFFFF'),
    (COL['c'],   'c[t]\nconsumo entrada',       C['gr_osc'], 'FFFFFF'),
    (COL['t1d'], 'a1 · vd_prev\n0.95 × vd ant.',C['vd_osc'], 'FFFFFF'),
    (COL['t2d'], 'b1 · vs_prev\n=0 (b1=0 por diseno)',C['vd_osc'], 'FFFFFF'),
    (COL['t3d'], '(1−a1) · c\n0.05 × entrada', C['vd_osc'], 'FFFFFF'),
    (COL['t4d'], '−g1 · sd_prev\n−0.5 × spike_d_prev',C['vd_osc'],'FFFFFF'),
    (COL['vd'],  '= vd[t]\nSUMA 4 terminos',   '145A32',    'FFFFFF'),
    (COL['sd'],  'sd[t]\nspike dendrita',       '1A5276',    'FFFFFF'),
    (COL['t1s'], 'a2 · vs_prev\n0.05 × vs ant.',C['rj_osc'], 'FFFFFF'),
    (COL['t2s'], 'b2 · vd[t]\n−0.9 × vd nuevo',C['rj_osc'], 'FFFFFF'),
    (COL['t3s'], '(1−a2) · c\n0.95 × entrada', C['rj_osc'], 'FFFFFF'),
    (COL['t4s'], '−g2 · ss_prev\n−1.0 × spike_s_prev',C['rj_osc'],'FFFFFF'),
    (COL['vs'],  '= vs[t]\nSUMA 4 terminos',   '922B21',    'FFFFFF'),
    (COL['ss'],  'ss[t]\nspike soma',           '1A5276',    'FFFFFF'),
    (COL['mix'], 's_mix[t]\nκ·sd+(1−κ)·ss',    C['mr_osc'], 'FFFFFF'),
]
for col, txt, bg, fg in HEADS:
    c = ws.cell(row=3, column=col, value=txt)
    c.font = Font(bold=True, color='FF'+fg, size=9)
    c.fill = fill(bg)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = brd()
ws.row_dimensions[3].height = 40

# Separadores (columnas I y P) — toda la altura
for sep_col in [COL['sep1'], COL['sep2']]:
    for r in range(1, N_ROWS + 10):
        ws.cell(row=r, column=sep_col).fill = fill('808080')

# ── Fila 4: ESTADO INICIAL (t = −1, todo cero) ───────────────────────────────
INIT_ROW = 4
ws.merge_cells(f'A{INIT_ROW}:B{INIT_ROW}')
c = ws.cell(row=INIT_ROW, column=1,
    value='Estado inicial  →  vd = 0  |  vs = 0  |  sd = 0  |  ss = 0')
c.font = Font(italic=True, size=9, color='FF555555')
c.fill = fill(C['gr_clr'])
c.alignment = Alignment(horizontal='center', vertical='center')
c.border = brd()

# Solo escribir los 4 valores de estado inicial necesarios para las primeras formulas
for col, val, fmt in [(COL['vd'],0,  '0.0000'),   # G4 = vd inicial
                       (COL['sd'],0,  '0'),         # H4 = sd inicial
                       (COL['vs'],0,  '0.0000'),    # N4 = vs inicial
                       (COL['ss'],0,  '0')]:        # O4 = ss inicial
    c = ws.cell(row=INIT_ROW, column=col, value=val)
    c.font = Font(bold=True, size=11, color='FF555555', italic=True)
    c.fill = fill(C['gr_clr'])
    c.number_format = fmt
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = brd()
ws.row_dimensions[INIT_ROW].height = 18

# ── Filas 5 … 5+167: UNA FILA POR HORA con formulas Excel reales ─────────────
for i in range(N_ROWS):
    row  = INIT_ROW + 1 + i      # fila Excel de este paso
    prev = row - 1                # fila Excel del paso anterior
    d_r  = DATA_ROW_START + i    # fila en hoja Datos para c[t=i]

    bg = C['gr_clr'] if i % 2 == 0 else C['blanco']

    # t
    c = ws.cell(row=row, column=COL['t'], value=i)
    c.font = Font(bold=True, size=10, color='FF1F3864')
    c.fill = fill(C['am'])
    c.number_format = '0'
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = brd()

    # c[t] — viene de la hoja Datos
    put(ws, row, COL['c'],
        formula=f"=Datos!$B${d_r}",
        bg=C['am'], bold=True, fmt='0.000000')

    # ── TERMINOS DENDRITA ─────────────────────────────────────────────────────
    # t1 = a1 * vd_prev  (vd del paso anterior esta en col G, fila prev)
    put(ws, row, COL['t1d'],
        formula=f"={PA1}*{L['vd']}{prev}",
        bg=C['vd_clr'], fmt='0.000000')

    # t2 = b1 * vs_prev  (vs del paso anterior esta en col N, fila prev)
    put(ws, row, COL['t2d'],
        formula=f"={PB1}*{L['vs']}{prev}",
        bg=C['vd_clr'], fmt='0.000000')

    # t3 = (1 - a1) * c[t]
    put(ws, row, COL['t3d'],
        formula=f"=(1-{PA1})*{L['c']}{row}",
        bg=C['vd_clr'], fmt='0.000000')

    # t4 = -g1 * sd_prev  (sd del paso anterior esta en col H, fila prev)
    put(ws, row, COL['t4d'],
        formula=f"=-{PG1}*{L['sd']}{prev}",
        bg=C['vd_clr'], fmt='0.000000')

    # vd[t] = suma de los 4 terminos
    c = ws.cell(row=row, column=COL['vd'])
    c.value = (f"={L['t1d']}{row}+{L['t2d']}{row}"
               f"+{L['t3d']}{row}+{L['t4d']}{row}")
    c.font = Font(bold=True, size=11, color='FF145A32')
    c.fill = fill(C['vd_brt'])
    c.number_format = '0.0000'
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = brd(thick_right=True)

    # sd[t] = IF(vd >= v_th, 1, 0)
    put(ws, row, COL['sd'],
        formula=f"=IF({L['vd']}{row}>={PVTH},1,0)",
        bg=C['am'], bold=True, fmt='0')

    # ── TERMINOS SOMA ─────────────────────────────────────────────────────────
    # t5 = a2 * vs_prev
    put(ws, row, COL['t1s'],
        formula=f"={PA2}*{L['vs']}{prev}",
        bg=C['rj_clr'], fmt='0.000000')

    # t6 = b2 * vd[t]  ← usa el vd recien calculado en esta misma fila
    put(ws, row, COL['t2s'],
        formula=f"={PB2}*{L['vd']}{row}",
        bg=C['rj_clr'], fmt='0.000000')

    # t7 = (1 - a2) * c[t]
    put(ws, row, COL['t3s'],
        formula=f"=(1-{PA2})*{L['c']}{row}",
        bg=C['rj_clr'], fmt='0.000000')

    # t8 = -g2 * ss_prev
    put(ws, row, COL['t4s'],
        formula=f"=-{PG2}*{L['ss']}{prev}",
        bg=C['rj_clr'], fmt='0.000000')

    # vs[t] = suma de los 4 terminos
    c = ws.cell(row=row, column=COL['vs'])
    c.value = (f"={L['t1s']}{row}+{L['t2s']}{row}"
               f"+{L['t3s']}{row}+{L['t4s']}{row}")
    c.font = Font(bold=True, size=11, color='FF7B0000')
    c.fill = fill(C['rj_brt'])
    c.number_format = '0.0000'
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = brd(thick_right=True)

    # ss[t] = IF(vs >= v_th, 1, 0)
    put(ws, row, COL['ss'],
        formula=f"=IF({L['vs']}{row}>={PVTH},1,0)",
        bg=C['am'], bold=True, fmt='0')

    # ── SALIDA ────────────────────────────────────────────────────────────────
    # s_mix = kappa*sd + (1-kappa)*ss
    put(ws, row, COL['mix'],
        formula=f"={PKAP}*{L['sd']}{row}+(1-{PKAP})*{L['ss']}{row}",
        bg=C['mr_clr'], bold=True, fg='4A235A', fmt='0.00')

    ws.row_dimensions[row].height = 17

# Fila de totales: contar spikes
tot = INIT_ROW + 1 + N_ROWS
sim_first = INIT_ROW + 1
sim_last  = tot

for col, label, formula, bg_c in [
    (COL['t'],   'TOTALES', None,                                C['gr_osc']),
    (COL['sd'],  None,      f"=COUNTIF({L['sd']}{sim_first}:{L['sd']}{sim_last},1)", '1A5276'),
    (COL['ss'],  None,      f"=COUNTIF({L['ss']}{sim_first}:{L['ss']}{sim_last},1)", '1A5276'),
    (COL['mix'], None,      f"=COUNTIF({L['mix']}{sim_first}:{L['mix']}{sim_last},\">0\")",C['mr_osc']),
]:
    c = ws.cell(row=tot+1, column=col)
    c.value = label if label else formula
    c.font = Font(bold=True, size=10, color='FFFFFFFF')
    c.fill = fill(bg_c)
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.number_format = '0'
    c.border = brd()

ws.merge_cells(f'A{tot+1}:B{tot+1}')
c = ws.cell(row=tot+1, column=1, value='TOTAL DISPAROS →')
c.font = Font(bold=True, size=10, color='FFFFFFFF')
c.fill = fill(C['gr_osc'])
c.alignment = Alignment(horizontal='right', vertical='center')
ws.row_dimensions[tot+1].height = 18

# Fila explicacion zeros
leg = tot + 2
ws.merge_cells(f'A{leg}:Q{leg}')
c = ws.cell(row=leg, column=1,
    value='NOTAS: (1) Columna D "b1·vs_prev" es siempre 0 porque b1=0 por diseno del paper '
          '(el soma no retroalimenta a la dendrita en la configuracion base). '
          '(2) Columna F "-g1·sd_prev" y col M "-g2·ss_prev" son 0 cuando no hubo spike en el paso anterior — '
          'el reset solo actua si la neurona disparo. '
          '(3) Con v_th=0.5 y datos [0,1] la dendrita dispara cuando vd supera el umbral; '
          'el soma (vs) puede no disparar si la inhibicion b2 lo mantiene bajo — es comportamiento correcto.')
c.font = Font(italic=True, size=9, color='FF222222')
c.fill = fill('FFF3CD')
c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws.row_dimensions[leg].height = 40

leg2 = leg + 1
ws.merge_cells(f'A{leg2}:Q{leg2}')
c = ws.cell(row=leg2, column=1,
    value='PARA EXPERIMENTAR: cambia v_th en Parametros!B10 '
          '(mas bajo = mas disparos)  |  cambia b1 en B6 a 0.1 para activar retroalimentacion soma→dendrita  '
          '|  cambia b2 en B7 a -0.3 para que el soma dispare mas.')
c.font = Font(italic=True, size=9, color='FF1A5276')
c.fill = fill('D6E4F0')
c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws.row_dimensions[leg2].height = 28

# Formato condicional: sd=1 y ss=1 en verde intenso, s_mix>0 en morado
spike_fill  = PatternFill(bgColor='00B050')
mix_fill    = PatternFill(bgColor='9933CC')
spike_font  = Font(bold=True, color='FFFFFFFF')
mix_font    = Font(bold=True, color='FFFFFFFF')

sd_range  = f"{L['sd']}{INIT_ROW+1}:{L['sd']}{tot}"
ss_range  = f"{L['ss']}{INIT_ROW+1}:{L['ss']}{tot}"
mix_range = f"{L['mix']}{INIT_ROW+1}:{L['mix']}{tot}"

for rng in [sd_range, ss_range]:
    ws.conditional_formatting.add(rng, Rule(
        type='cellIs', operator='equal', formula=['1'],
        dxf=DifferentialStyle(fill=spike_fill, font=spike_font)
    ))
ws.conditional_formatting.add(mix_range, Rule(
    type='cellIs', operator='greaterThan', formula=['0'],
    dxf=DifferentialStyle(fill=mix_fill, font=mix_font)
))

ws.freeze_panes = 'C5'

# =============================================================================
# HOJA 4 — GRAFICO: c[t] vs vd[t] vs vs[t]
# =============================================================================
ws_g = wb.create_sheet('Grafico')
ws_g.column_dimensions['A'].width = 6
ws_g.column_dimensions['B'].width = 14
ws_g.column_dimensions['C'].width = 14
ws_g.column_dimensions['D'].width = 14

# Cabeceras
for col, txt in [(1,'t'),(2,'c[t] entrada'),(3,'vd[t] dendrita'),(4,'vs[t] soma')]:
    put(ws_g, 1, col, txt, bg=C['gr_osc'], fg='FFFFFF', bold=True)

# Formulas que referencian la hoja Simulacion
for i in range(N_ROWS):
    sim_row = INIT_ROW + 1 + i
    ws_g.cell(row=2+i, column=1, value=i).number_format = '0'
    ws_g.cell(row=2+i, column=2).value = f"=Simulacion!{L['c']}{sim_row}"
    ws_g.cell(row=2+i, column=3).value = f"=Simulacion!{L['vd']}{sim_row}"
    ws_g.cell(row=2+i, column=4).value = f"=Simulacion!{L['vs']}{sim_row}"
    for col in range(1, 5):
        ws_g.cell(row=2+i, column=col).number_format = '0.0000'

# Grafico de lineas
chart = LineChart()
chart.title  = 'TS-LIF — c[t] entrada  |  vd[t] dendrita (paso-bajo)  |  vs[t] soma (paso-alto)'
chart.style  = 10
chart.y_axis.title = 'Valor'
chart.x_axis.title = 'Hora t  (0 = Lun 00:00  →  167 = Dom 23:00)'
chart.height = 15
chart.width  = 32

for col_idx, color, width in [(2,'1F3864',10000),(3,'1E5631',18000),(4,'7B0000',18000)]:
    ref = Reference(ws_g, min_col=col_idx, min_row=1, max_row=N_ROWS+1)
    chart.add_data(ref, titles_from_data=True)
    chart.series[-1].graphicalProperties.line.solidFill = color
    chart.series[-1].graphicalProperties.line.width = width

cats = Reference(ws_g, min_col=1, min_row=2, max_row=N_ROWS+1)
chart.set_categories(cats)
ws_g.add_chart(chart, 'F2')

# =============================================================================
wb.save(OUT)
print(f"Excel guardado: {OUT}")
print()
print("  Hoja 'Parametros'  — B4:B11 editables + estabilidad con formulas IF/SQRT")
print("  Hoja 'Datos'       — 168 pasos UCI Electricity, c[t=0] en B3")
print("  Hoja 'Simulacion'  — 168 filas con formulas encadenadas, sin valores hardcodeados")
print("  Hoja 'Grafico'     — c[t] vs vd[t] vs vs[t] en 168 horas")
print()
print("  Referencias correctas en Simulacion:")
print(f"    a1  = {PA1}")
print(f"    a2  = {PA2}")
print(f"    b2  = {PB2}")
print(f"    v_th= {PVTH}")
print(f"    kap = {PKAP}")
